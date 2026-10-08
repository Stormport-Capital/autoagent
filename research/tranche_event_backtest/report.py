"""Step 5: the Excel workbook. Every net-R figure is a live formula off the
Assumptions tab, so changing slippage or borrow there recalculates every
trade, summary row and entry-day row.

    write_workbook(path, trades, events, excluded, checks)
"""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


STRATEGIES = ("5/10 EMA cross", "10/20 EMA cross", "VWAP fail (Russo Trigger B)", "Opening fade")
SHORT_ONLY = {"VWAP fail (Russo Trigger B)", "Opening fade"}   # engine.py:543
TIMEFRAMES = ("1h", "15m", "5m")
DIRECTIONS = ("long", "short")
EVENT_TYPES = ("INTRADAY_100", "CLOSE_100")
SMALL_SAMPLE = 20
BOLD = Font(bold=True)
INPUT = PatternFill("solid", fgColor="FFF2CC")

# Trades tab layout: (header, key or None for a formula / helper)
TRADE_COLS = [
    ("Event ID", "event_id"), ("Symbol", "symbol"), ("Event type", "event_type"),
    ("Event date", "event_date"), ("HELD/FADED", "held_faded"), ("Strategy", "strategy"),
    ("Timeframe", "timeframe"), ("Direction", "direction"), ("Signal time", "signal_time"),
    ("Entry day", "entry_day"), ("Entry time", "entry_time"), ("Entry price", "entry_price"),
    ("Stop", "stop"), ("Exit time", "exit_time"), ("Exit price", "exit_price"),
    ("Exit reason", "exit_reason"), ("Holding period (trading days)", "holding_days"),
    ("Gapped stop", "gapped"), ("Status", "status"), ("Gross R", "gross_r"), ("Net R", None),
    ("Also under event", "also_under"), ("Size (shares)", "qty"),
    ("Leg 1 exit time", "l1t"), ("Leg 1 exit price", "l1p"), ("Leg 1 size", "l1q"),
    ("Leg 2 exit time", "l2t"), ("Leg 2 exit price", "l2p"), ("Leg 2 size", "l2q"),
    ("Splits rescaled in window", "splits_applied"),
    ("Risk basis (helper)", "ema_unit"), ("Risk per share, gross (helper)", "risk_per_share"),
    ("Leg 1 borrow base (helper)", "l1b"), ("Leg 2 borrow base (helper)", "l2b"),
    ("Group (helper)", None), ("Losing streak (helper)", None),
]
COL = {h: get_column_letter(i + 1) for i, (h, _k) in enumerate(TRADE_COLS)}


def _naive(v):
    """Excel has no time zones: write New York wall-clock time."""
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    return v


def _flat(r: dict) -> dict:
    legs = r["legs"]
    last = legs[-1]
    f = dict(r)
    f["exit_time"] = last["time"]
    f["exit_price"] = last["price"]
    f["gapped"] = "yes" if r["gapped_stop"] else "no"
    for n, leg in ((1, legs[0]), (2, legs[1] if len(legs) > 1 else None)):
        f[f"l{n}t"] = leg["time"] if leg else None
        f[f"l{n}p"] = leg["price"] if leg else None
        f[f"l{n}q"] = leg["qty"] if leg else None
        f[f"l{n}b"] = leg["borrow_base"] if leg else None
    return f


def _sort_key(r):
    return (r["event_type"], r["strategy"], r["timeframe"], r["direction"],
            r["status"] != "CLOSED", r["entry_time"])


def assumptions_sheet(wb):
    ws = wb.active
    ws.title = "Assumptions"
    rows = [
        ("Assumption", "Value", "Notes"),
        ("Slippage per side (%)", 0.05, "Engine default 5 bps each side (store.py:18). Applied adversely to every entry and exit fill, as the engine does (engine.py:672-673, 713-719). Set to 0 for gross."),
        ("Short borrow rate (% per year)", 10.0, "Engine default 10%/yr (store.py:13)."),
        ("Borrow day count", 360, "Engine default 360 (store.py:14): rate / day count per calendar night held, on the prior close (engine.py:730-740). Shorts only; Friday to Monday is 3 nights."),
        ("Borrow per calendar night (%)", "=B3/B4", "Derived; do not edit."),
    ]
    for r in rows:
        ws.append(r)
    for c in ("B2", "B3", "B4"):
        ws[c].fill = INPUT
    ws["A1"].font = ws["B1"].font = ws["C1"].font = BOLD
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["C"].width = 110
    ws.append(())
    ws.append(("Gross R comes from the engine run with zero slippage. Net R on every trade is a live formula off the yellow cells.",))
    ws.append(("Net R = (P&L per share after slippage - borrow per share) / stop distance after slippage, the engine's own R (engine.py:819-820).",))
    ws.append(("Trades that scaled out (VWAP half at the 10-day MA, rest at the 20-day) are one row with both legs; the engine splits shares half/half, so with costs its own net R can differ from this row by under 0.01R.",))


def trades_sheet(wb, rows):
    ws = wb.create_sheet("Trades")
    ws.append([h for h, _k in TRADE_COLS])
    for c in ws[1]:
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    rows = sorted(rows, key=_sort_key)
    c = COL
    for i, r in enumerate(rows, start=2):
        f = _flat(r)
        vals = []
        for h, k in TRADE_COLS:
            vals.append(_naive(f.get(k)) if k else None)
        ws.append(vals)
        s, sign = "Assumptions!$B$2/100", f'IF({c["Direction"]}{i}="long",1,-1)'
        ef = f'{c["Entry price"]}{i}*(1+{sign}*{s})'
        q = f'{c["Size (shares)"]}{i}'
        leg = lambda n: (f'N({c[f"Leg {n} size"]}{i})/{q}*{sign}*'
                         f'(N({c[f"Leg {n} exit price"]}{i})*(1-{sign}*{s})-{ef})')
        bor = (f'Assumptions!$B$3/100/Assumptions!$B$4*('
               f'N({c["Leg 1 size"]}{i})/{q}*N({c["Leg 1 borrow base (helper)"]}{i})+'
               f'N({c["Leg 2 size"]}{i})/{q}*N({c["Leg 2 borrow base (helper)"]}{i}))')
        risk = (f'IF({c["Risk basis (helper)"]}{i}="stop",ABS({c["Stop"]}{i}-{ef}),'
                f'IF({c["Risk basis (helper)"]}{i}="pct",{c["Risk per share, gross (helper)"]}{i}/{c["Entry price"]}{i}*{ef},'
                f'{c["Risk per share, gross (helper)"]}{i}))')
        ws[f'{c["Net R"]}{i}'] = f"=({leg(1)}+{leg(2)}-{bor})/{risk}"
        ws[f'{c["Group (helper)"]}{i}'] = (f'={c["Event type"]}{i}&"|"&{c["Strategy"]}{i}&"|"&'
                                            f'{c["Timeframe"]}{i}&"|"&{c["Direction"]}{i}')
        loss = f'AND({c["Status"]}{i}="CLOSED",{c["Net R"]}{i}<=0)'
        if i == 2:
            ws[f'{c["Losing streak (helper)"]}{i}'] = f"=IF({loss},1,0)"
        else:
            same = f'{c["Group (helper)"]}{i}={c["Group (helper)"]}{i - 1}'
            ws[f'{c["Losing streak (helper)"]}{i}'] = (
                f'=IF({loss},IF({same},{c["Losing streak (helper)"]}{i - 1}+1,1),0)')
    for col in (c["Signal time"], c["Entry time"], c["Exit time"], c["Leg 1 exit time"], c["Leg 2 exit time"]):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
    for col in (c["Event date"],):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd"
    for col in (c["Gross R"], c["Net R"]):
        for cell in ws[col][1:]:
            cell.number_format = "0.00"
    ws.freeze_panes = "C2"
    return len(rows) + 1


def _crit(n, **kw):
    """COUNTIFS-style criteria pairs over Trades rows 2..n."""
    m = {"event_type": "Event type", "strategy": "Strategy", "timeframe": "Timeframe",
         "direction": "Direction", "status": "Status", "held_faded": "HELD/FADED", "entry_day": "Entry day"}
    parts = []
    for k, v in kw.items():
        col = COL[m[k]]
        parts.append(f"Trades!${col}$2:${col}${n},{v}")
    return ",".join(parts)


def _rng(name, n):
    col = COL[name]
    return f"Trades!${col}$2:${col}${n}"


def summary_rows(ws, n, keys: list[dict], extra_cols: list[str]):
    head = extra_cols + ["Closed trades", "Open trades", "Win %", "Avg gross R", "Total gross R",
                         "Avg net R", "Total net R", "Avg holding (trading days)",
                         "Largest losing streak", "Worst trade (net R)", "Best trade (net R)", "Note"]
    ws.append(head)
    for cell in ws[ws.max_row]:
        cell.font = BOLD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for k in keys:
        r = ws.max_row + 1
        q = {key: f'"{v}"' if isinstance(v, str) else v for key, v in k["crit"].items()}
        closed = _crit(n, **q, status='"CLOSED"')
        opened = _crit(n, **q, status='"OPEN"')
        nx = len(extra_cols)
        L = lambda j: get_column_letter(nx + j)
        cells = list(k["labels"]) + [
            f"=COUNTIFS({closed})",
            f"=COUNTIFS({opened})",
            f'=IF({L(1)}{r}=0,"",COUNTIFS({closed},{_rng("Net R", n)},">0")/{L(1)}{r})',
            f'=IF({L(1)}{r}=0,"",AVERAGEIFS({_rng("Gross R", n)},{closed}))',
            f'=SUMIFS({_rng("Gross R", n)},{closed})',
            f'=IF({L(1)}{r}=0,"",AVERAGEIFS({_rng("Net R", n)},{closed}))',
            f'=SUMIFS({_rng("Net R", n)},{closed})',
            f'=IF({L(1)}{r}=0,"",AVERAGEIFS({_rng("Holding period (trading days)", n)},{closed}))',
            f'=IF({L(1)}{r}=0,"",_xlfn.MAXIFS({_rng("Losing streak (helper)", n)},{closed}))',
            f'=IF({L(1)}{r}=0,"",_xlfn.MINIFS({_rng("Net R", n)},{closed}))',
            f'=IF({L(1)}{r}=0,"",_xlfn.MAXIFS({_rng("Net R", n)},{closed}))',
            k.get("note") or f'=IF({L(1)}{r}<{SMALL_SAMPLE},"small sample","")',
        ]
        ws.append(cells)
        for j, fmt in ((3, "0%"), (4, "0.00"), (5, "0.00"), (6, "0.00"), (7, "0.00"), (8, "0.0"),
                       (10, "0.00"), (11, "0.00")):
            ws.cell(row=r, column=nx + j).number_format = fmt


def _keys(event_type, held_faded=None):
    out = []
    for s in STRATEGIES:
        for tf in TIMEFRAMES:
            for d in DIRECTIONS:
                crit = {"event_type": event_type, "strategy": s, "timeframe": tf, "direction": d}
                labels = [s, tf, d]
                if held_faded:
                    crit["held_faded"] = held_faded
                    labels = [held_faded] + labels
                note = "rule is short-only (engine.py:543)" if s in SHORT_ONLY and d == "long" else None
                out.append({"crit": crit, "labels": labels, "note": note})
    return out


def write_workbook(path, trade_rows, events, excluded, checks):
    wb = Workbook()
    assumptions_sheet(wb)
    n = trades_sheet(wb, trade_rows)
    for et in EVENT_TYPES:
        ws = wb.create_sheet(f"Summary {et}")
        summary_rows(ws, n, _keys(et), ["Strategy", "Timeframe", "Direction"])
        ws.freeze_panes = "D2"
    ws = wb.create_sheet("HELD vs FADED")
    summary_rows(ws, n, _keys("INTRADAY_100", "HELD") + _keys("INTRADAY_100", "FADED"),
                 ["HELD/FADED", "Strategy", "Timeframe", "Direction"])
    ws = wb.create_sheet("By entry day")
    ws.append(["Event type", "Strategy", "Timeframe", "Direction", "Entry day", "Closed trades",
               "Win %", "Avg net R", "Total net R"])
    for cell in ws[1]:
        cell.font = BOLD
    for et in EVENT_TYPES:
        for s in STRATEGIES:
            for tf in TIMEFRAMES:
                for d in DIRECTIONS:
                    if s in SHORT_ONLY and d == "long":
                        continue
                    for day in range(0, 26):
                        r = ws.max_row + 1
                        closed = _crit(n, event_type=f'"{et}"', strategy=f'"{s}"', timeframe=f'"{tf}"',
                                       direction=f'"{d}"', entry_day=day, status='"CLOSED"')
                        ws.append([et, s, tf, d, day, f"=COUNTIFS({closed})",
                                   f'=IF(F{r}=0,"",COUNTIFS({closed},{_rng("Net R", n)},">0")/F{r})',
                                   f'=IF(F{r}=0,"",AVERAGEIFS({_rng("Net R", n)},{closed}))',
                                   f'=SUMIFS({_rng("Net R", n)},{closed})'])
                        ws.cell(row=r, column=7).number_format = "0%"
                        ws.cell(row=r, column=8).number_format = "0.00"
                        ws.cell(row=r, column=9).number_format = "0.00"
    ws.freeze_panes = "F2"
    for title, rows in (("Events", events), ("Excluded", excluded)):
        ws = wb.create_sheet(title)
        if rows:
            head = list(dict.fromkeys(k for r in rows for k in r))   # union of columns, first-seen order
            ws.append(head)
            for cell in ws[1]:
                cell.font = BOLD
            for r in rows:
                vals = [r.get(k) for k in head]
                ws.append([("yes" if v else "no") if isinstance(v, bool) else
                           float(v) if isinstance(v, Decimal) else _naive(v) for v in vals])
            ws.freeze_panes = "B2"
    ws = wb.create_sheet("Checks")
    for line in checks:
        ws.append(line if isinstance(line, (list, tuple)) else [line])
    ws.column_dimensions["A"].width = 140
    wb.save(path)
